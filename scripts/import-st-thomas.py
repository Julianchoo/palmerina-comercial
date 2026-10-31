#!/usr/bin/env python3
"""Prepare a PRIVATE seed and plan from the original CAD PDF and lot spreadsheet.

Dependencies: Python 3.10+, pymupdf, openpyxl, numpy, scipy.
Example (paths and output must remain private):
  python scripts/import-st-thomas.py --plan /private/plan.pdf \
    --spreadsheet /private/lots.xlsx --output .private-st-thomas

Until sold IDs are supplied and remaining statuses explicitly confirmed, every
lot starts unavailable. Confirm using --sold-lots N,N,N --remaining-status
available. Existing output records keep their commercial fields; this script
never connects to a database. A database seed importer must INSERT ON CONFLICT
DO NOTHING, never replace an already edited inventory.

Geometry comes from the PDF's Lote_Poligono CAD strokes, including Bezier curves.
It does not use guessed rectangles or label-distance Voronoi regions. Curves are
flattened with maximum control-point deviation 0.005 PDF point; CAD endpoint
round-off is snapped within 0.1 PDF point. The extracted planar graph is split
at touching parcel dividers, then its bounded faces are matched to lot labels.
"""
from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re
import uuid

import fitz
import numpy as np
from openpyxl import load_workbook
from scipy.spatial import cKDTree

SNAP = 0.1
CURVE_ERROR = 0.005
MUTABLE_FIELDS = ("id", "street", "areaM2", "priceUsd", "status", "notes", "updatedAt", "version")


def polygon_area(points):
    return sum(a[0] * b[1] - b[0] * a[1]
               for a, b in zip(points, points[1:] + points[:1])) / 2


def contains(point, polygon):
    x, y = point
    inside = False
    for a, b in zip(polygon, polygon[1:] + polygon[:1]):
        if ((a[1] > y) != (b[1] > y)
                and x < (b[0] - a[0]) * (y - a[1]) / (b[1] - a[1]) + a[0]):
            inside = not inside
    return inside


def flatten_bezier(a, b, c, end, result):
    chord = end - a
    length = float(np.linalg.norm(chord))

    def deviation(point):
        offset = point - a
        return (abs(chord[0] * offset[1] - chord[1] * offset[0]) / length
                if length else float(np.linalg.norm(offset)))

    if max(deviation(b), deviation(c)) <= CURVE_ERROR:
        result.append((tuple(a), tuple(end)))
        return
    ab, bc, ce = (a + b) / 2, (b + c) / 2, (c + end) / 2
    abc, bce = (ab + bc) / 2, (bc + ce) / 2
    midpoint = (abc + bce) / 2
    flatten_bezier(a, ab, abc, midpoint, result)
    flatten_bezier(midpoint, bce, ce, end, result)


def parcel_faces(page):
    segments = []
    for drawing in page.get_drawings():
        if drawing.get("layer") != "Lote_Poligono":
            continue
        for item in drawing["items"]:
            if item[0] == "l":
                segments.append((tuple(item[1]), tuple(item[2])))
            elif item[0] == "c":
                flatten_bezier(*(np.array(tuple(p)) for p in item[1:]), segments)
            else:
                raise ValueError(f"Unsupported parcel CAD item: {item[0]}")
    if not segments:
        raise ValueError("The PDF has no Lote_Poligono CAD boundaries")

    points = np.array([point for segment in segments for point in segment])
    parent = list(range(len(points)))

    def root(index):
        while parent[index] != index:
            parent[index] = parent[parent[index]]
            index = parent[index]
        return index

    for a, b in cKDTree(points).query_pairs(SNAP):
        parent[root(b)] = root(a)
    groups = defaultdict(list)
    for index in range(len(points)):
        groups[root(index)].append(index)
    coordinates = {key: tuple(points[members].mean(axis=0))
                   for key, members in groups.items()}
    edges = {tuple(sorted((root(2 * i), root(2 * i + 1))))
             for i in range(len(segments)) if root(2 * i) != root(2 * i + 1)}

    # A parcel divider may touch the middle of a longer CAD stroke. Split the
    # stroke at that real endpoint before walking the planar graph's faces.
    keys = list(coordinates)
    positions = np.array([coordinates[key] for key in keys])
    spatial_index = cKDTree(positions)
    split_edges = set()
    for a, b in edges:
        av, bv = np.array(coordinates[a]), np.array(coordinates[b])
        delta = bv - av
        length2 = float(delta @ delta)
        candidates = spatial_index.query_ball_point((av + bv) / 2, math.sqrt(length2) / 2 + SNAP)
        on_stroke = [(0.0, a), (1.0, b)]
        for index in candidates:
            key = keys[index]
            if key in (a, b):
                continue
            projection = float((positions[index] - av) @ delta / length2)
            if (0.00001 < projection < 0.99999
                    and np.linalg.norm(positions[index] - (av + projection * delta)) < SNAP):
                on_stroke.append((projection, key))
        on_stroke.sort()
        for (_, u), (_, v) in zip(on_stroke, on_stroke[1:]):
            if u != v:
                split_edges.add(tuple(sorted((u, v))))

    neighbours = {key: set() for key in coordinates}
    for a, b in split_edges:
        neighbours[a].add(b)
        neighbours[b].add(a)
    ordered = {key: sorted(values, key=lambda other: math.atan2(
        coordinates[other][1] - coordinates[key][1],
        coordinates[other][0] - coordinates[key][0])) for key, values in neighbours.items()}
    visited = set()
    faces = []
    for a, b in split_edges:
        for start in ((a, b), (b, a)):
            if start in visited:
                continue
            face, current = [], start
            while current not in visited:
                visited.add(current)
                u, v = current
                face.append(coordinates[u])
                options = ordered[v]
                current = (v, options[(options.index(u) - 1) % len(options)])
            if current != start:
                raise ValueError("CAD boundary walk did not form a closed face")
            if polygon_area(face) > 1:
                # Retain CAD curve shape, round only much below snap tolerance.
                faces.append([[round(x, 4), round(y, 4)] for x, y in face])
    return faces


def text_spans(page):
    return [span for block in page.get_text("dict")["blocks"]
            for line in block.get("lines", []) for span in line["spans"]]


def center(span):
    x0, y0, x1, y1 = span["bbox"]
    return ((x0 + x1) / 2, (y0 + y1) / 2)


def spreadsheet_lots(path):
    book = load_workbook(path, data_only=True)
    sheet = book.active
    lots = []
    # Each CALLE heading is followed by adjacent LOTE / M2 columns. Read only
    # the block above its subtotal; do not mistake subtotal counts for IDs.
    for column in range(1, sheet.max_column):
        heading = sheet.cell(4, column).value
        if not isinstance(heading, str) or not heading.startswith("CALLE "):
            continue
        group = []
        for row in range(6, sheet.max_row + 1):
            number, area = sheet.cell(row, column).value, sheet.cell(row, column + 1).value
            if not isinstance(number, int) or not isinstance(area, (int, float)):
                continue
            if area >= 10000:
                if number != len(group) or abs(float(area) - sum(x["areaM2"] for x in group)) > .02:
                    raise ValueError(f"Subtotal does not reconcile for {heading}")
                break
            if not math.isfinite(area) or area <= 0:
                raise ValueError(f"Invalid surface for lot {number}")
            group.append({"number": number, "street": heading, "areaM2": float(area)})
        lots.extend(group)
    if not lots or len({lot["number"] for lot in lots}) != len(lots):
        raise ValueError("No lots or duplicate lot numbers in spreadsheet")
    declared_count, declared_area = sheet["H1"].value, sheet["H2"].value
    if len(lots) != declared_count or abs(sum(lot["areaM2"] for lot in lots) - declared_area) > .02:
        raise ValueError("Inventory does not reconcile with spreadsheet totals")
    return sorted(lots, key=lambda lot: lot["number"])


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--plan", type=Path, required=True)
    parser.add_argument("--spreadsheet", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=Path(".private-st-thomas"))
    parser.add_argument("--sold-lots", default="", help="Comma-separated confirmed sold lot numbers")
    parser.add_argument("--remaining-status", choices=("available", "unavailable"), default="unavailable")
    parser.add_argument("--expected-sold-count", type=int, default=3)
    args = parser.parse_args()
    output = args.output.resolve()
    # Do not accidentally export private sources into the site's static tree.
    if "public" in output.parts or output == Path.cwd().resolve():
        parser.error("Use a dedicated PRIVATE output directory, never public/ or the repository root")
    output.mkdir(parents=True, exist_ok=True)
    document = fitz.open(args.plan)
    if len(document) != 1:
        parser.error("Expected a single-page site plan")
    page = document[0]
    faces = parcel_faces(page)
    spans = text_spans(page)
    lots = spreadsheet_lots(args.spreadsheet)
    sold_numbers = {int(value.strip()) for value in args.sold_lots.split(",") if value.strip()}
    if sold_numbers - {lot["number"] for lot in lots}:
        parser.error("Sold IDs must belong to the spreadsheet inventory")
    confirmed = len(sold_numbers) == args.expected_sold_count
    if args.remaining_status == "available" and not confirmed:
        parser.error("Confirm all expected sold lot IDs before making remaining lots available")
    warnings = []
    if not confirmed:
        warnings.append("Pendiente identificar los 3 lotes vendidos; inventario sin confirmar, no ofrecer como disponible.")
    existing_file = output / "inventory.json"
    existing = json.loads(existing_file.read_text()) if existing_file.exists() else {}
    existing_lots = {lot["number"]: lot for lot in existing.get("lots", [])}
    now = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    audit_records, geometric_ratios = [], []
    for lot in lots:
        number = lot["number"]
        labels = [span for span in spans if span["text"].strip() == str(number)]
        if len(labels) != 1:
            raise ValueError(f"Lot {number} has {len(labels)} PDF labels")
        matches = [face for face in faces if contains(center(labels[0]), face)]
        if len(matches) != 1:
            raise ValueError(f"Lot {number} has {len(matches)} containing CAD parcel faces")
        polygon = matches[0]
        printed_areas = [span["text"] for span in spans
                         if re.fullmatch(r"[\d.,]+m²", span["text"].strip())
                         and contains(center(span), polygon)]
        if len(printed_areas) != 1:
            raise ValueError(f"Lot {number}: could not uniquely validate printed PDF surface")
        pdf_area = float(printed_areas[0].split("m")[0].replace(",", "."))
        if abs(pdf_area - lot["areaM2"]) > .011:
            warnings.append(f"Lote {number}: superficie Excel {lot['areaM2']} m² / PDF {pdf_area} m².")
        geometry_area = polygon_area(polygon)
        geometric_ratios.append(geometry_area / lot["areaM2"])
        lot.update({
            "id": str(uuid.uuid5(uuid.NAMESPACE_URL, f"st-thomas:lot:{number}")),
            "priceUsd": 108000 if lot["areaM2"] == 760 else 168000 if lot["areaM2"] >= 1400 else None,
            "status": "sold" if number in sold_numbers else args.remaining_status,
            "notes": "" if confirmed else "Estado comercial pendiente de confirmar.",
            "polygon": polygon,
            "updatedAt": now,
            "version": 1,
        })
        previous = existing_lots.get(number)
        if previous:
            for field in MUTABLE_FIELDS:
                lot[field] = previous[field]
        audit_records.append({"number": number, "spreadsheetAreaM2": lot["areaM2"],
                              "pdfAreaM2": pdf_area, "geometryAreaPdfPoints": geometry_area,
                              "label": list(center(labels[0])), "faceVertexCount": len(polygon)})
    median_ratio = float(np.median(geometric_ratios))
    for record in audit_records:
        scaled_area = record["geometryAreaPdfPoints"] / median_ratio
        record["geometryAreaM2AtMedianScale"] = round(scaled_area, 3)
        discrepancy = abs(scaled_area - record["pdfAreaM2"]) / record["pdfAreaM2"]
        if discrepancy > .005:
            warnings.append(f"Lote {record['number']}: área dibujada aproximada {scaled_area:.2f} m² "
                            f"frente a superficie declarada {record['pdfAreaM2']:.2f} m²; "
                            "se conserva superficie declarada y límite CAD real.")
    if existing_lots:
        if sold_numbers:
            confirmed = confirmed and all(lot["status"] == "sold" for lot in lots
                                          if lot["number"] in sold_numbers)
        else:
            confirmed = bool(existing.get("sourceAudit", {}).get("statusesConfirmed", False))
    if existing_lots and sold_numbers:
        warnings.append("Se conservaron los estados ya existentes; revisar ventas confirmadas en el editor.")
    seed = {"schemaVersion": 1,
            "plan": {"width": page.rect.width, "height": page.rect.height, "imageFile": "plan.png"},
            "lots": lots,
            "sourceAudit": {"lotCount": len(lots),
                            "totalAreaM2": round(sum(lot["areaM2"] for lot in lots), 2),
                            "statusesConfirmed": confirmed, "warnings": warnings}}
    # Deterministic seed records and preservation of existing mutable fields make
    # re-running the source preparation safe for an existing private seed.
    (output / "inventory.json").write_text(json.dumps(seed, ensure_ascii=False, indent=2) + "\n")
    page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).save(str(output / "plan.png"))
    source_audit = {"sourcePdfSha256": hashlib.sha256(args.plan.read_bytes()).hexdigest(),
                    "sourceExcelSha256": hashlib.sha256(args.spreadsheet.read_bytes()).hexdigest(),
                    "cadLayer": "Lote_Poligono", "cadFaces": len(faces),
                    "curveFlatnessPdfPoints": CURVE_ERROR, "snapTolerancePdfPoints": SNAP,
                    "areaRatioPdfPointsPerM2": median_ratio, "warnings": warnings,
                    "records": audit_records}
    (output / "audit.json").write_text(json.dumps(source_audit, ensure_ascii=False, indent=2) + "\n")
    print(f"Prepared {len(lots)} private lots; {len(faces)} CAD parcel faces; "
          f"{seed['sourceAudit']['totalAreaM2']:.2f} m². Statuses confirmed: {confirmed}.")
    for warning in warnings:
        print(warning)


if __name__ == "__main__":
    main()
